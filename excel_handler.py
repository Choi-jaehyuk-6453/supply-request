import os
import shutil
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook, Workbook
from openpyxl.utils.dataframe import dataframe_to_rows


DB_FILE = "management_db.xlsx"
MASTER_SHEET = "Master_Data"
SUMMARY_SHEET = "Monthly_Summary"

NETWORK_PATH = r"\\192.168.0.100\보안팀\4.(미래_다원) 2026년 피복(장비), 경비용품비 사용현황"
NETWORK_FILE = os.path.join(NETWORK_PATH, "management_db.xlsx") if os.path.exists(NETWORK_PATH) else None


def sync_to_network():
    """네트워크 경로로 엑셀 파일 동기화"""
    try:
        if NETWORK_PATH and os.path.exists(NETWORK_PATH):
            shutil.copy2(DB_FILE, NETWORK_FILE)
            return True, "네트워크 저장소에 동기화 완료"
    except PermissionError:
        return False, "네트워크 저장소 접근 권한이 없습니다."
    except Exception as e:
        return False, f"네트워크 동기화 오류: {str(e)}"
    return False, "네트워크 경로를 사용할 수 없습니다. 로컬 파일에만 저장됩니다."


def initialize_excel():
    """엑셀 파일 초기화"""
    if not os.path.exists(DB_FILE):
        wb = Workbook()
        
        ws_master = wb.active
        ws_master.title = MASTER_SHEET
        ws_master.append(['날짜', '구분', '법인명', '현장명', '신청자', '품목명', '규격', '수량', '단가', '합계금액'])
        
        ws_summary = wb.create_sheet(SUMMARY_SHEET)
        ws_summary.append(['년월', '법인명', '현장명', '총액'])
        
        wb.save(DB_FILE)
        wb.close()


def append_to_master(data, app_type, items):
    """Master_Data 시트에 데이터 추가"""
    initialize_excel()
    
    try:
        wb = load_workbook(DB_FILE)
        ws = wb[MASTER_SHEET]
        
        date_str = data.get('application_date', datetime.now().strftime('%Y-%m-%d'))
        company = data.get('company', '')
        site_name = data.get('site_name', '')
        applicant = data.get('applicant', '')
        
        for item in items:
            if app_type == '피복':
                product_name = item.get('product', '')
                spec = f"{item.get('top_size', '')}/{item.get('bottom_size', '')}"
                quantity = 1
                unit_price = item.get('unit_price', 0)
            else:
                product_name = item.get('name', '')
                spec = item.get('spec', '')
                quantity = item.get('quantity', 0)
                unit_price = item.get('unit_price', 0)
            
            total = quantity * unit_price
            
            ws.append([
                date_str,
                app_type,
                company,
                site_name,
                applicant,
                product_name,
                spec,
                quantity,
                unit_price,
                total
            ])
        
        wb.save(DB_FILE)
        wb.close()
        
        update_monthly_summary()
        
        network_success, network_msg = sync_to_network()
        if network_success:
            return True, "데이터가 성공적으로 저장되었습니다. (네트워크 동기화 완료)"
        else:
            return True, f"로컬에 저장되었습니다. ({network_msg})"
    
    except PermissionError:
        return False, "엑셀 파일이 열려있어 저장할 수 없습니다. 파일을 닫고 다시 시도해주세요."
    except Exception as e:
        return False, f"데이터 저장 중 오류가 발생했습니다: {str(e)}"


def update_monthly_summary():
    """Monthly_Summary 시트 갱신"""
    try:
        df = pd.read_excel(DB_FILE, sheet_name=MASTER_SHEET)
        
        if df.empty:
            return
        
        df['날짜'] = pd.to_datetime(df['날짜'], errors='coerce')
        df['년월'] = df['날짜'].dt.strftime('%Y-%m')
        
        summary = df.groupby(['년월', '법인명', '현장명'])['합계금액'].sum().reset_index()
        summary.columns = ['년월', '법인명', '현장명', '총액']
        
        wb = load_workbook(DB_FILE)
        
        if SUMMARY_SHEET in wb.sheetnames:
            del wb[SUMMARY_SHEET]
        
        ws_summary = wb.create_sheet(SUMMARY_SHEET)
        
        for r_idx, row in enumerate(dataframe_to_rows(summary, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                ws_summary.cell(row=r_idx, column=c_idx, value=value)
        
        wb.save(DB_FILE)
        wb.close()
        
    except Exception as e:
        print(f"월별 집계 갱신 오류: {str(e)}")


def get_master_data():
    """Master_Data 조회"""
    initialize_excel()
    try:
        df = pd.read_excel(DB_FILE, sheet_name=MASTER_SHEET)
        return df
    except Exception:
        return pd.DataFrame()


def get_monthly_summary():
    """Monthly_Summary 조회"""
    initialize_excel()
    try:
        df = pd.read_excel(DB_FILE, sheet_name=SUMMARY_SHEET)
        return df
    except Exception:
        return pd.DataFrame()


def load_price_list():
    """단가표 로드 (별도 엑셀 파일에서)"""
    price_file = "price_list.xlsx"
    if os.path.exists(price_file):
        try:
            df = pd.read_excel(price_file)
            return df
        except Exception:
            pass
    
    default_prices = {
        '경비물품': {
            '커피모카골드 100T': 15000,
            '화장지(흰들베이직 50) 50M 무포': 25000,
            '3M슈퍼그립200(겨울용)': 8000,
        },
        '피복': {
            '회색동복': 50000,
            '하복': 40000,
            '점퍼': 60000,
        }
    }
    return default_prices
